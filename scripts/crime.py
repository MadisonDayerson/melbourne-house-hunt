"""Extract suburb-level crime counts and the metro suburb list from the CSA workbook.

Needs openpyxl. Outputs:
  data/crime_suburb_jun2026.json  {lga: {...}, suburb: {name: counts}, lga_of: {name: [lgas]}}
  data/suburbs.json               [{suburb, postcode, lga, region}] for every metro suburb
"""
import openpyxl, json, collections, pathlib

D = pathlib.Path(__file__).resolve().parent.parent / "data"

# The 31 metropolitan councils, grouped into the regions used in the app's filter
REGIONS = {
    "Inner city": ["Melbourne", "Yarra", "Port Phillip", "Stonnington"],
    "North": ["Merri-bek", "Darebin", "Banyule", "Moonee Valley", "Hume", "Whittlesea", "Nillumbik"],
    "West": ["Maribyrnong", "Hobsons Bay", "Brimbank", "Wyndham", "Melton"],
    "East": ["Boroondara", "Manningham", "Whitehorse", "Maroondah", "Knox", "Yarra Ranges"],
    "South-east": ["Monash", "Glen Eira", "Greater Dandenong", "Casey", "Cardinia"],
    "Bayside & peninsula": ["Bayside", "Kingston", "Frankston", "Mornington Peninsula"],
}
REGION_OF = {lga: r for r, lgas in REGIONS.items() for lga in lgas}

wb = openpyxl.load_workbook(D / "csa_lga_offences_jun2026.xlsx", read_only=True)
lga_rate = {}
for r in wb["Table 01"].iter_rows(min_row=2, values_only=True):
    if r[0] == 2026 and r[3] and isinstance(r[5], (int, float)):
        lga_rate[r[3].strip()] = (r[4], round(r[5]))
missing = [l for l in REGION_OF if l not in lga_rate]
assert not missing, f"LGA names not found in workbook: {missing}"

sub = collections.defaultdict(collections.Counter)
lgas = collections.defaultdict(collections.Counter)   # suburb -> lga -> offences
pcs = collections.defaultdict(collections.Counter)    # suburb -> postcode -> offences
for r in wb["Table 03"].iter_rows(min_row=2, values_only=True):
    if r[0] != 2026 or not r[2] or r[2].strip() not in REGION_OF:
        continue
    lga, pc, name, sg, n = r[2].strip(), r[3], r[4].strip(), r[7] or "", r[8] or 0
    c = sub[name]
    c["total"] += n
    if sg.startswith("B3"): c["burglary"] += n
    if sg.startswith("B4"): c["theft"] += n
    if sg.startswith("A2"): c["assault"] += n
    if sg.startswith("A5"): c["robbery"] += n
    lgas[name][lga] += n
    pcs[name][pc] += n

json.dump({"lga": {k: lga_rate[k] for k in REGION_OF}, "suburb": sub,
           "lga_of": {k: [l for l, _ in v.most_common()] for k, v in lgas.items()}},
          open(D / "crime_suburb_jun2026.json", "w"), indent=1)

suburbs = []
for name in sorted(sub):
    lga = lgas[name].most_common(1)[0][0]
    suburbs.append({"suburb": name, "postcode": str(pcs[name].most_common(1)[0][0]), "lga": lga, "region": REGION_OF[lga]})
json.dump(suburbs, open(D / "suburbs.json", "w"), indent=1)
print(len(suburbs), "suburbs;", collections.Counter(s["region"] for s in suburbs))
